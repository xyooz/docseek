from __future__ import annotations

import argparse
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from docseek.chunks import DocumentChunk, _iter_xlsx_chunks, _xlsx_chunk_content


@dataclass(frozen=True, slots=True)
class BackendResult:
    name: str
    seconds: float
    chunks: tuple[DocumentChunk, ...]


def build_workbook(path: Path, *, rows: int, sheets: int, columns: int) -> None:
    workbook = Workbook(write_only=True)
    for sheet_index in range(sheets):
        sheet = workbook.create_sheet(title=f"客户明细{sheet_index + 1}")
        sheet.append([f"字段{i + 1}" for i in range(columns)])
        for row_index in range(1, rows + 1):
            row: list[object] = []
            for column in range(columns):
                if column == 0:
                    value: object = f"客户{sheet_index:02d}-{row_index:07d}"
                elif column == 1:
                    value = "客户经理" if row_index % 7 == 0 else "普通客户"
                elif column == 2:
                    value = "身份证有效期" if row_index % 113 == 0 else "资料完整"
                elif column == 3:
                    value = row_index * 100 + sheet_index
                elif column == 4 and row_index % 5 == 0:
                    value = None
                elif column == 5:
                    value = row_index / 10
                else:
                    value = f"业务字段{column}-{row_index % 97}"
                row.append(value)
            sheet.append(row)
    workbook.save(path)


def _calamine_cell_text(value: object) -> str:
    """Match openpyxl's stable search-text representation for numeric cells.

    python-calamine exposes Excel numbers as floats, including integral values
    such as ``100.0``. openpyxl returns those same integral cells as ``100`` in
    our current production path. Preserve the existing DocSeek index text while
    benchmarking the faster backend instead of treating a Python type detail as
    a search-semantic change.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _calamine_row_text(row: tuple[object, ...]) -> str:
    cells = [_calamine_cell_text(value) for value in row]
    while cells and not cells[-1]:
        cells.pop()
    if not cells or not any(cells):
        return ""
    return "\t".join(cells)


def iter_calamine_chunks(path: Path, *, rows_per_chunk: int = 200):
    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_path(path)
    ordinal = 0
    try:
        for sheet_name in workbook.sheet_names:
            sheet = workbook.get_sheet_by_name(sheet_name)
            buffer: list[str] = []
            first_row = 1
            last_row = 0
            for row_no, row in enumerate(sheet.iter_rows(), start=1):
                row_text = _calamine_row_text(tuple(row))
                if row_text:
                    if not buffer:
                        first_row = row_no
                    buffer.append(row_text)
                    last_row = row_no
                if buffer and len(buffer) >= rows_per_chunk:
                    yield DocumentChunk(
                        ordinal,
                        f"工作表 {sheet_name} · 行 {first_row}-{last_row}",
                        _xlsx_chunk_content(sheet_name, buffer),
                    )
                    ordinal += 1
                    buffer = []
            if buffer:
                yield DocumentChunk(
                    ordinal,
                    f"工作表 {sheet_name} · 行 {first_row}-{last_row}",
                    _xlsx_chunk_content(sheet_name, buffer),
                )
                ordinal += 1
    finally:
        workbook.close()


def measure(name: str, iterator_factory) -> BackendResult:
    started = time.perf_counter()
    chunks = tuple(iterator_factory())
    return BackendResult(name, time.perf_counter() - started, chunks)


def first_difference(
    left: tuple[DocumentChunk, ...],
    right: tuple[DocumentChunk, ...],
) -> str | None:
    if len(left) != len(right):
        return f"chunk count differs: {len(left)} != {len(right)}"
    for index, (a, b) in enumerate(zip(left, right, strict=True)):
        if a.location != b.location:
            return f"chunk {index} location differs: {a.location!r} != {b.location!r}"
        if a.content != b.content:
            limit = min(len(a.content), len(b.content))
            offset = next((i for i in range(limit) if a.content[i] != b.content[i]), limit)
            return (
                f"chunk {index} content differs at char {offset}: "
                f"openpyxl={a.content[offset:offset + 80]!r} "
                f"calamine={b.content[offset:offset + 80]!r}"
            )
    return None


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(
        description="Compare DocSeek XLSX chunk extraction through openpyxl and python-calamine"
    )
    parser.add_argument("--rows", type=int, default=20_000)
    parser.add_argument("--sheets", type=int, default=2)
    parser.add_argument("--columns", type=int, default=10)
    parser.add_argument("--rows-per-chunk", type=int, default=200)
    parser.add_argument(
        "--require-equivalent",
        action="store_true",
        help="fail when the two backends do not produce identical DocSeek chunks",
    )
    args = parser.parse_args()

    with tempfile.TemporaryDirectory(prefix="docseek-xlsx-backend-") as temp_dir:
        path = Path(temp_dir) / "网点客户明细.xlsx"
        build_workbook(
            path,
            rows=max(1, args.rows),
            sheets=max(1, args.sheets),
            columns=max(6, args.columns),
        )

        openpyxl_result = measure(
            "openpyxl",
            lambda: _iter_xlsx_chunks(
                path,
                rows_per_chunk=max(1, args.rows_per_chunk),
            ),
        )
        calamine_result = measure(
            "calamine",
            lambda: iter_calamine_chunks(
                path,
                rows_per_chunk=max(1, args.rows_per_chunk),
            ),
        )
        difference = first_difference(openpyxl_result.chunks, calamine_result.chunks)
        speedup = (
            openpyxl_result.seconds / calamine_result.seconds
            if calamine_result.seconds > 0
            else float("inf")
        )

        print("DocSeek XLSX backend comparison")
        print(
            f"rows/sheet={args.rows:,} sheets={args.sheets} columns={args.columns} "
            f"xlsx={path.stat().st_size / 1024 / 1024:.2f}MiB"
        )
        print(
            f"openpyxl={openpyxl_result.seconds:.3f}s "
            f"chunks={len(openpyxl_result.chunks):,}"
        )
        print(
            f"calamine={calamine_result.seconds:.3f}s "
            f"chunks={len(calamine_result.chunks):,} speedup={speedup:.2f}x"
        )
        print("equivalent=yes" if difference is None else f"equivalent=no · {difference}")

        if difference is not None and args.require_equivalent:
            return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
