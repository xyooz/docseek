from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from .document_types import KNOWN_DOCUMENT_EXTENSIONS


# Compatibility export used by the watcher/indexer. This now means “known
# document formats worth attempting locally”, not “every format has a built-in
# Python parser”. Missing optional adapters are surfaced as persistent index
# issues instead of being silently ignored.
SUPPORTED_EXTENSIONS = set(KNOWN_DOCUMENT_EXTENSIONS)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".log", ".csv", ".tsv"}:
        return _extract_plain_text(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)

    # Legacy API compatibility for newly registered optional formats. Keep the
    # same extraction broker as the chunk indexer instead of implementing a
    # second set of format parsers here.
    if suffix in SUPPORTED_EXTENSIONS:
        from .chunks import iter_document_chunks

        return "\n".join(chunk.content for chunk in iter_document_chunks(path))

    raise ValueError(f"Unsupported file type: {suffix}")


def _extract_plain_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding, errors="strict")
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_docx(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            chunks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def _extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            chunks.append(f"[Sheet: {worksheet.title}]")
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    chunks.append("\t".join(values))
    finally:
        workbook.close()
    return "\n".join(chunks)


def _extract_pptx(path: Path) -> str:
    presentation = Presentation(path)
    chunks: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_chunks: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_chunks.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    slide_chunks.append("\t".join(cell.text for cell in row.cells))
        if slide_chunks:
            chunks.append(f"[Slide: {index}]")
            chunks.extend(slide_chunks)
    return "\n".join(chunks)


def _extract_pdf(path: Path) -> str:
    chunks: list[str] = []
    with pymupdf.open(path) as document:
        for page_no, page in enumerate(document, start=1):
            text = page.get_text("text")
            if text:
                chunks.append(f"[Page: {page_no}]\n{text}")
    return "\n".join(chunks)
