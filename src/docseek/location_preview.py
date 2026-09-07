"""Bounded, read-only source previews. Worker output stays in memory."""
from __future__ import annotations

import base64
import html
import json
import re
import sys
from pathlib import Path


def preview_kind(extension: str, location: str) -> str | None:
    if extension.lower() == ".pdf" and re.fullmatch(r"第\s*\d+\s*页", location):
        return "pdf"
    if extension.lower() == ".xlsx" and re.fullmatch(r"工作表 (.+?) · 行 (\d+)-(\d+)", location):
        return "xlsx"
    return None


def render_preview(path: str, location: str, modified_time: float, size: int) -> dict:
    source = Path(path)
    before = source.stat()
    if before.st_size != size or abs(before.st_mtime - modified_time) > 0.001:
        raise ValueError("文件已变化，请等待索引更新后重新搜索。")
    kind = preview_kind(source.suffix, location)
    if kind == "pdf":
        import pymupdf

        page_no = int(re.search(r"\d+", location).group())
        with pymupdf.open(source) as document:
            if document.needs_pass:
                raise ValueError("此 PDF 需要密码，请使用原程序打开。")
            if not 1 <= page_no <= len(document):
                raise ValueError("页码已失效，请刷新索引。")
            page = document[page_no - 1]
            scale = min(1.5, 1400 / max(page.rect.width, page.rect.height, 1))
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
            result = {"html": f"<h3>{html.escape(location)}</h3><p>原文件页面预览</p><img src='preview:page'>",
                      "png": base64.b64encode(pixmap.tobytes("png")).decode("ascii")}
    elif kind == "xlsx":
        from openpyxl import load_workbook

        match = re.fullmatch(r"工作表 (.+?) · 行 (\d+)-(\d+)", location)
        sheet, start, end = match.groups()
        start, end = int(start), int(end)
        if start < 1 or end < start:
            raise ValueError("行位置无效，请刷新索引。")
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet]
            # Keep the first row as context, without guessing that it is a header.
            first = next(worksheet.iter_rows(min_row=1, max_row=1, max_col=50, values_only=True), ()) if start > 1 else ()
            reference = ""
            if first:
                reference = "<p>第 1 行参考（不一定是表头）：" + html.escape(
                    " | ".join("" if cell is None else str(cell)[:100] for cell in first).rstrip(" | ")) + "</p>"
            # Retain real row numbers including blanks.
            stop = min(end, start + 199)
            rows = []
            for number, values in enumerate(worksheet.iter_rows(min_row=start, max_row=stop, max_col=50, values_only=True), start):
                cells = ["" if value is None else str(value)[:1000] for value in values]
                while cells and not cells[-1]:
                    cells.pop()
                rows.append((number, cells))
            width = max((len(cells) for _, cells in rows), default=0)
            from openpyxl.utils import get_column_letter
            table = "<table border='1' cellspacing='0' cellpadding='4'><tr><th>行</th>"
            table += "".join(f"<th>{get_column_letter(i + 1)}</th>" for i in range(width)) + "</tr>"
            for number, cells in rows:
                table += f"<tr><th>{number}</th>" + "".join(
                    f"<td>{html.escape(cells[i] if i < len(cells) else '')}</td>" for i in range(width)) + "</tr>"
            result = {"html": f"<h3>{html.escape(location)}</h3><p>当前显示行 {start}-{stop}，最多 200 行、前 50 列；单元格最多 1000 字符，公式显示文件中保存的缓存值。</p>" + reference + table + "</table>"}
        finally:
            workbook.close()
    else:
        raise ValueError("此格式暂不支持原文件位置预览，请查看命中上下文。")
    after = source.stat()
    if (before.st_mtime_ns, before.st_size) != (after.st_mtime_ns, after.st_size):
        raise ValueError("预览期间文件已变化，请重新搜索。")
    return result


def main(argv=None) -> int:
    args = sys.argv[1:] if argv is None else argv
    import os
    try:
        result = render_preview(args[0], args[1], float(args[2]), int(args[3]))
    except Exception as exc:
        result = {"error": str(exc)}
    payload = json.dumps(result, ensure_ascii=True).encode("ascii")
    if sys.platform == "win32":
        # Windowed executables have no Python stdout / CRT fd 1. QProcess
        # supplies a Win32 pipe handle, which remains available in that case.
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.GetStdHandle.argtypes = [wintypes.DWORD]
        kernel.GetStdHandle.restype = wintypes.HANDLE
        kernel.WriteFile.argtypes = [wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD,
                                     ctypes.POINTER(wintypes.DWORD), ctypes.c_void_p]
        kernel.WriteFile.restype = wintypes.BOOL
        handle = kernel.GetStdHandle(-11 & 0xFFFFFFFF)
        offset = 0
        while offset < len(payload):
            data = payload[offset:offset + 65536]
            written = wintypes.DWORD()
            if not kernel.WriteFile(handle, data, len(data), ctypes.byref(written), None) or not written.value:
                return 1
            offset += written.value
    else:
        with os.fdopen(os.dup(1), "wb") as output:
            output.write(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
