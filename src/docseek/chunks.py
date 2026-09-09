from __future__ import annotations

import io
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Iterable, Iterator

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation


@dataclass(slots=True, frozen=True)
class DocumentChunk:
    ordinal: int
    location: str
    content: str


TEXT_EXTENSIONS = {".txt", ".md", ".log", ".csv", ".tsv"}
HTML_EXTENSIONS = {".html", ".htm", ".xhtml"}
TEXT_ENCODING_SAMPLE_BYTES = 65_536
TEXT_ENCODING_CANDIDATES = ("utf-8", "gb18030")
XLSX_PROGRESS_ROW_INTERVAL = 1_000
ChunkProgressCallback = Callable[[str, int], None]


def iter_document_chunks(
    path: Path,
    *,
    target_chars: int = 12_000,
    xlsx_rows_per_chunk: int = 200,
    on_progress: ChunkProgressCallback | None = None,
    cancelled: Callable[[], bool] | None = None,
    persistent_worker=None,
) -> Iterator[DocumentChunk]:
    """Public extraction entrypoint used by the production indexer.

    Every supported format now passes through the capability-aware broker and
    streaming DocIR layer before being converted back to the existing v7
    ``DocumentChunk`` contract. The compatibility conversion is lossless, so
    this architectural upgrade does not change FTS text or require a schema
    migration. Structural labels may become richer as extractors learn titles.
    """
    from .extraction_broker import DEFAULT_EXTRACTION_BROKER

    kwargs = {
        "target_chars": target_chars,
        "spreadsheet_rows_per_chunk": xlsx_rows_per_chunk,
        "on_progress": on_progress,
    }
    if cancelled is not None:
        kwargs["cancelled"] = cancelled
    if persistent_worker is not None:
        kwargs["persistent_worker"] = persistent_worker
    yield from DEFAULT_EXTRACTION_BROKER.iter_chunks(path, **kwargs)


def _iter_direct_document_chunks(
    path: Path,
    *,
    target_chars: int = 12_000,
    spreadsheet_rows_per_chunk: int = 200,
    on_progress: ChunkProgressCallback | None = None,
) -> Iterator[DocumentChunk]:
    """Mature built-in parsers used only by ``DirectDocumentAdapter``.

    Keeping this function separate from the public broker entrypoint prevents
    adapter recursion while preserving the already-tested low-memory streaming
    implementations.
    """
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        yield from _iter_text_chunks(path, target_chars=target_chars)
    elif suffix in HTML_EXTENSIONS:
        yield from _iter_html_chunks(path, target_chars=target_chars)
    elif suffix == ".xml":
        yield from _iter_xml_chunks(path, target_chars=target_chars)
    elif suffix == ".docx":
        yield from _iter_docx_chunks(path, target_chars=target_chars)
    elif suffix == ".xlsx":
        yield from _iter_xlsx_chunks(
            path,
            rows_per_chunk=spreadsheet_rows_per_chunk,
            on_progress=on_progress,
        )
    elif suffix == ".pptx":
        yield from _iter_pptx_chunks(path)
    elif suffix == ".pdf":
        yield from _iter_pdf_chunks(path)
    else:
        raise ValueError(f"Unsupported direct file type: {suffix}")


def _iter_text_chunks(path: Path, *, target_chars: int) -> Iterator[DocumentChunk]:
    """Decode and chunk plain text while keeping the large-file path streaming.

    The first bounded read is already enough to decode a small file. When that
    normalized text also fits in one target chunk, construct the chunk directly
    instead of routing it through StringIO, a line generator, a temporary list
    and a final join. Larger small files still use the exact line-based chunking
    contract, while files above the probe limit rewind the same handle and
    stream through TextIOWrapper as before.
    """
    with path.open("rb") as raw:
        probe = raw.read(TEXT_ENCODING_SAMPLE_BYTES + 1)
        has_more = len(probe) > TEXT_ENCODING_SAMPLE_BYTES

        if not has_more:
            _encoding, decoded = _decode_text_bytes(probe)
            # Match TextIOWrapper/StringIO universal-newline behavior used by
            # the previous implementation before applying chunk boundaries.
            normalized = decoded.replace("\r\n", "\n").replace("\r", "\n")
            if len(normalized) < target_chars:
                content = normalized.strip()
                if content:
                    line_count = normalized.count("\n")
                    if not normalized.endswith("\n"):
                        line_count += 1
                    yield DocumentChunk(0, f"行 1-{max(1, line_count)}", content)
                return

            yield from _iter_text_chunks_from_lines(
                normalized.splitlines(keepends=True),
                target_chars=target_chars,
            )
            return

        sample = probe[:TEXT_ENCODING_SAMPLE_BYTES]
        encoding = _detect_text_encoding_sample(sample)
        raw.seek(0)
        with io.TextIOWrapper(
            raw,
            encoding=encoding,
            errors="ignore",
            newline=None,
        ) as text:
            yield from _iter_text_chunks_from_lines(text, target_chars=target_chars)


def _iter_text_chunks_from_lines(
    lines: Iterable[str],
    *,
    target_chars: int,
) -> Iterator[DocumentChunk]:
    ordinal = 0
    start_line = 1
    line_no = 0
    buffer: list[str] = []
    char_count = 0

    for line_no, line in enumerate(lines, start=1):
        buffer.append(line.rstrip("\n"))
        char_count += len(line)
        if char_count >= target_chars:
            content = "\n".join(buffer).strip()
            if content:
                yield DocumentChunk(ordinal, f"行 {start_line}-{line_no}", content)
                ordinal += 1
            buffer = []
            char_count = 0
            start_line = line_no + 1

    content = "\n".join(buffer).strip()
    if content:
        end_line = max(start_line, line_no)
        yield DocumentChunk(ordinal, f"行 {start_line}-{end_line}", content)


def _decode_text_bytes(data: bytes) -> tuple[str, str]:
    for encoding in TEXT_ENCODING_CANDIDATES:
        try:
            return encoding, data.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
    return "utf-8", data.decode("utf-8", errors="ignore")


def _detect_text_encoding_sample(sample: bytes) -> str:
    encoding, _decoded = _decode_text_bytes(sample)
    return encoding


class _VisibleHtmlParser(HTMLParser):
    """Collect visible HTML text while dropping executable and style content."""

    _BLOCK_TAGS = frozenset(
        {"address", "article", "aside", "blockquote", "br", "dd", "div", "dl",
         "dt", "figcaption", "figure", "footer", "form", "h1", "h2", "h3",
         "h4", "h5", "h6", "header", "hr", "li", "main", "nav", "ol", "p",
         "pre", "section", "table", "td", "th", "title", "tr", "ul"}
    )
    _SKIP_TAGS = frozenset({"script", "style", "noscript", "svg", "canvas"})
    _HEADING_TAGS = frozenset({"title", "h1", "h2", "h3", "h4", "h5", "h6"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[tuple[str, str]] = []
        self._parts: list[str] = []
        self._skip_depth = 0
        self._heading_tag = ""
        self._heading_parts: list[str] = []
        self._current_heading = ""

    def handle_starttag(self, tag: str, attrs) -> None:
        del attrs
        tag = tag.casefold()
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag in self._HEADING_TAGS:
            self._flush_line()
            self._heading_tag = tag
            self._heading_parts = []
        elif tag in self._BLOCK_TAGS:
            self._flush_line()

    def handle_startendtag(self, tag: str, attrs) -> None:
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if tag in self._SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if tag == self._heading_tag:
            heading = _clean_location_title(" ".join(self._heading_parts))
            self._flush_line(heading=heading)
            if heading:
                self._current_heading = heading
            self._heading_tag = ""
            self._heading_parts = []
        elif tag in self._BLOCK_TAGS:
            self._flush_line()

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = " ".join(data.split())
        if not text:
            return
        self._parts.append(text)
        if self._heading_tag:
            self._heading_parts.append(text)

    def close(self) -> None:
        super().close()
        self._flush_line()

    def _flush_line(self, *, heading: str = "") -> None:
        text = " ".join(self._parts).strip()
        self._parts = []
        if text:
            self.lines.append((heading or self._current_heading, text))


def _iter_html_chunks(path: Path, *, target_chars: int) -> Iterator[DocumentChunk]:
    """Extract visible HTML locally without paying one Tika process per file."""
    with path.open("rb") as handle:
        sample = handle.read(TEXT_ENCODING_SAMPLE_BYTES)
        encoding = _detect_text_encoding_sample(sample)
        handle.seek(0)
        text = handle.read().decode(encoding, errors="ignore")

    parser = _VisibleHtmlParser()
    parser.feed(text)
    parser.close()

    ordinal = 0
    buffer: list[str] = []
    char_count = 0
    chunk_heading = ""
    for heading, line in parser.lines:
        if buffer and char_count >= target_chars:
            location = f"网页内容 {ordinal + 1}"
            if chunk_heading:
                location += f" · 标题 {chunk_heading}"
            yield DocumentChunk(ordinal, location, "\n".join(buffer))
            ordinal += 1
            buffer = []
            char_count = 0
            chunk_heading = ""
        if heading and line == heading:
            # Prefer the most specific heading encountered in this chunk over
            # an earlier document <title> label.
            chunk_heading = heading
        elif not buffer:
            chunk_heading = heading
        buffer.append(line)
        char_count += len(line)

    if buffer:
        location = f"网页内容 {ordinal + 1}"
        if chunk_heading:
            location += f" · 标题 {chunk_heading}"
        yield DocumentChunk(ordinal, location, "\n".join(buffer))


class _TolerantXmlTextParser(HTMLParser):
    """Collect searchable XML text without requiring a single strict root.

    Real developer and office trees contain XML fragments, vendor entities and
    occasionally truncated metadata files. HTMLParser gives us a bounded,
    non-executing markup tokenizer that still extracts useful text from those
    files instead of classifying every non-canonical document as corrupt.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=False)
        self.elements: list[str] = []
        self.lines: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        element = tag.rsplit(":", 1)[-1][:80]
        self.elements.append(element)
        values = [" ".join(str(value).split()) for _name, value in attrs if value]
        text = " ".join(value for value in values if value)
        if text:
            self.lines.append((element, text))

    def handle_startendtag(self, tag: str, attrs) -> None:
        element = tag.rsplit(":", 1)[-1][:80]
        values = [" ".join(str(value).split()) for _name, value in attrs if value]
        text = " ".join(value for value in values if value)
        if text:
            self.lines.append((element, text))

    def handle_endtag(self, tag: str) -> None:
        target = tag.rsplit(":", 1)[-1]
        for index in range(len(self.elements) - 1, -1, -1):
            if self.elements[index].casefold() == target.casefold():
                del self.elements[index:]
                return

    def handle_data(self, data: str) -> None:
        text = " ".join(data.split())
        if text:
            self.lines.append((self.elements[-1] if self.elements else "", text))

    def handle_entityref(self, name: str) -> None:
        # Keep custom entity names searchable without attempting expansion.
        if name:
            self.lines.append((self.elements[-1] if self.elements else "", name))

    def handle_charref(self, name: str) -> None:
        if name:
            self.lines.append((self.elements[-1] if self.elements else "", name))

    def drain_lines(self) -> list[tuple[str, str]]:
        lines, self.lines = self.lines, []
        return lines


def _markup_encoding(sample: bytes) -> str:
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    declaration = sample[:512].decode("ascii", errors="ignore")
    match = re.search(
        r"<\?xml\b[^>]*\bencoding\s*=\s*(['\"])([^'\"]+)\1",
        declaration,
        flags=re.IGNORECASE,
    )
    if match:
        candidate = match.group(2).strip()
        if candidate:
            try:
                "".encode(candidate)
                return candidate
            except LookupError:
                pass
    return _detect_text_encoding_sample(sample)


def _iter_xml_chunks(path: Path, *, target_chars: int) -> Iterator[DocumentChunk]:
    """Stream useful text from strict XML, fragments and vendor XML variants."""
    parser = _TolerantXmlTextParser()
    ordinal = 0
    buffer: list[str] = []
    char_count = 0
    current_element = ""

    def ready_chunks() -> Iterator[DocumentChunk]:
        nonlocal ordinal, buffer, char_count, current_element
        for element, text in parser.drain_lines():
            current_element = element or current_element
            buffer.append(text)
            char_count += len(text)
            if char_count >= target_chars:
                location = f"XML 内容 {ordinal + 1}"
                if current_element:
                    location += f" · 元素 {current_element}"
                yield DocumentChunk(ordinal, location, "\n".join(buffer))
                ordinal += 1
                buffer = []
                char_count = 0

    with path.open("rb") as raw:
        sample = raw.read(TEXT_ENCODING_SAMPLE_BYTES)
        encoding = _markup_encoding(sample)
        raw.seek(0)
        with io.TextIOWrapper(raw, encoding=encoding, errors="ignore", newline=None) as text:
            while data := text.read(TEXT_ENCODING_SAMPLE_BYTES):
                parser.feed(data)
                yield from ready_chunks()
    parser.close()
    yield from ready_chunks()

    if buffer:
        location = f"XML 内容 {ordinal + 1}"
        if current_element:
            location += f" · 元素 {current_element}"
        yield DocumentChunk(ordinal, location, "\n".join(buffer))


def _clean_location_title(text: str, *, max_chars: int = 80) -> str:
    """Normalize a human-readable structural title for compact result labels."""
    normalized = " ".join(text.split()).strip()
    if len(normalized) <= max_chars:
        return normalized
    return normalized[: max_chars - 3].rstrip() + "..."


def _is_docx_heading_style(style_name: str) -> bool:
    style = style_name.strip()
    folded = style.casefold()
    return (
        folded.startswith("heading")
        or folded == "title"
        or style.startswith("标题")
    )


def _iter_docx_chunks(path: Path, *, target_chars: int) -> Iterator[DocumentChunk]:
    document = Document(path)
    ordinal = 0
    buffer: list[str] = []
    char_count = 0
    section_start = 1
    block_no = 0
    current_heading = ""
    buffer_heading = ""

    def maybe_emit(force: bool = False) -> DocumentChunk | None:
        nonlocal ordinal, buffer, char_count, section_start, buffer_heading
        if not buffer or (not force and char_count < target_chars):
            return None
        content = "\n".join(buffer).strip()
        if not content:
            buffer = []
            char_count = 0
            buffer_heading = ""
            return None
        location = f"文档块 {section_start}-{block_no}"
        if buffer_heading:
            location += f" · 标题 {buffer_heading}"
        chunk = DocumentChunk(ordinal, location, content)
        ordinal += 1
        buffer = []
        char_count = 0
        buffer_heading = ""
        section_start = block_no + 1
        return chunk

    for element in document.element.body.iterchildren():
        heading = False
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            texts = [paragraph.text.strip()]
            style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
            heading = _is_docx_heading_style(style_name)
        elif element.tag.endswith("}tbl"):
            texts = ("\t".join(cell.text for cell in row.cells).strip()
                     for row in Table(element, document).rows)
        else:
            continue
        for text in texts:
            if not text:
                continue
            if heading:
                chunk = maybe_emit(force=True)
                if chunk:
                    yield chunk
                current_heading = _clean_location_title(text)
            block_no += 1
            if not buffer:
                buffer_heading = current_heading
            buffer.append(text)
            char_count += len(text)
            chunk = maybe_emit()
            if chunk:
                yield chunk

    chunk = maybe_emit(force=True)
    if chunk:
        yield chunk


def _xlsx_row_text(row: tuple[object, ...]) -> str:
    """Render one spreadsheet row without shifting later columns to the left."""
    cells = ["" if value is None else str(value) for value in row]
    while cells and not cells[-1]:
        cells.pop()
    if not cells or not any(cells):
        return ""
    return "\t".join(cells)


def _xlsx_chunk_content(sheet_title: str, rows: list[str]) -> str:
    return f"工作表: {sheet_title}\n" + "\n".join(rows)


def _iter_xlsx_chunks(
    path: Path,
    *,
    rows_per_chunk: int,
    on_progress: ChunkProgressCallback | None = None,
) -> Iterator[DocumentChunk]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    ordinal = 0
    try:
        for worksheet in workbook.worksheets:
            buffer: list[str] = []
            first_row = 1
            last_row = 0
            processed_row = 0
            progress_label = f"工作表 {worksheet.title}"
            for row_no, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                processed_row = row_no
                row_text = _xlsx_row_text(row)
                if row_text:
                    if not buffer:
                        first_row = row_no
                    buffer.append(row_text)
                    last_row = row_no
                if buffer and len(buffer) >= rows_per_chunk:
                    yield DocumentChunk(
                        ordinal,
                        f"工作表 {worksheet.title} · 行 {first_row}-{last_row}",
                        _xlsx_chunk_content(worksheet.title, buffer),
                    )
                    ordinal += 1
                    buffer = []
                if on_progress and row_no % XLSX_PROGRESS_ROW_INTERVAL == 0:
                    on_progress(progress_label, row_no)
            if buffer:
                yield DocumentChunk(
                    ordinal,
                    f"工作表 {worksheet.title} · 行 {first_row}-{last_row}",
                    _xlsx_chunk_content(worksheet.title, buffer),
                )
                ordinal += 1
            if (
                on_progress
                and processed_row > 0
                and processed_row % XLSX_PROGRESS_ROW_INTERVAL != 0
            ):
                on_progress(progress_label, processed_row)
    finally:
        workbook.close()


def _iter_pptx_chunks(path: Path) -> Iterator[DocumentChunk]:
    presentation = Presentation(path)
    ordinal = 0
    for slide_no, slide in enumerate(presentation.slides, start=1):
        content: list[str] = []
        title_shape = getattr(slide.shapes, "title", None)
        title_text = _clean_location_title(str(getattr(title_shape, "text", "") or ""))
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                content.append(str(text))
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    content.append("\t".join(cell.text for cell in row.cells))
        text = "\n".join(content).strip()
        if text:
            location = f"幻灯片 {slide_no}"
            if title_text:
                location += f" · 标题 {title_text}"
            yield DocumentChunk(ordinal, location, text)
            ordinal += 1


def _iter_pdf_chunks(path: Path) -> Iterator[DocumentChunk]:
    ordinal = 0
    with pymupdf.open(path) as document:
        for page_no, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                yield DocumentChunk(ordinal, f"第 {page_no} 页", text)
                ordinal += 1
