from __future__ import annotations

"""Verify the shipped Windows executable's persistent extraction contract.

This is intentionally a standalone packaging check rather than a normal unit
test.  It talks to the extracted ``DocSeek.exe`` directly, then asks that same
binary to run a real DirectoryIndexer job.  A one-shot fallback cannot satisfy
the PID and protocol assertions here.
"""

import argparse
import json
import os
import pickle
import queue
import subprocess
import tempfile
import threading
from collections import deque
from pathlib import Path


def _json_line(message: dict[str, object]) -> bytes:
    return (json.dumps(message, ensure_ascii=True, separators=(",", ":")) + "\n").encode(
        "ascii"
    )


class _PersistentProtocol:
    def __init__(self, executable: Path, work_dir: Path) -> None:
        creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        self.process = subprocess.Popen(
            [str(executable), "--docseek-extract-worker", "--persistent"],
            cwd=str(work_dir),
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            creationflags=creationflags,
        )
        assert self.process.stdin is not None
        assert self.process.stdout is not None
        assert self.process.stderr is not None
        self._messages: queue.Queue[bytes | None] = queue.Queue()
        self._stderr_tail: deque[bytes] = deque(maxlen=32)
        self._stdout_thread = threading.Thread(
            target=self._read_stdout,
            name="portable-persistent-worker-stdout",
            daemon=True,
        )
        self._stderr_thread = threading.Thread(
            target=self._read_stderr,
            name="portable-persistent-worker-stderr",
            daemon=True,
        )
        self._stdout_thread.start()
        self._stderr_thread.start()

    def _read_stdout(self) -> None:
        assert self.process.stdout is not None
        try:
            for line in iter(self.process.stdout.readline, b""):
                self._messages.put(line)
        finally:
            self._messages.put(None)

    def _read_stderr(self) -> None:
        assert self.process.stderr is not None
        for line in iter(self.process.stderr.readline, b""):
            self._stderr_tail.append(line)

    def send(self, message: dict[str, object]) -> None:
        assert self.process.stdin is not None
        self.process.stdin.write(_json_line(message))
        self.process.stdin.flush()

    def receive(self, timeout: float = 30.0) -> dict[str, object]:
        try:
            line = self._messages.get(timeout=timeout)
        except queue.Empty as exc:
            raise RuntimeError("timed out waiting for frozen worker protocol") from exc
        if line is None:
            stderr = b" ".join(self._stderr_tail).decode("utf-8", errors="replace")
            raise RuntimeError(
                f"frozen worker closed its protocol stream; stderr={stderr[-1000:]!r}"
            )
        try:
            message = json.loads(line.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise RuntimeError(f"invalid frozen worker protocol line: {line!r}") from exc
        if not isinstance(message, dict):
            raise RuntimeError(f"frozen worker protocol message is not an object: {message!r}")
        return message

    def terminate(self) -> None:
        if self.process.poll() is None:
            self.process.terminate()
            try:
                self.process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                self.process.kill()
                self.process.wait(timeout=5)


def _read_spool(path: Path) -> tuple[list[tuple[int, str, str]], int]:
    records: list[tuple[int, str, str]] = []
    with path.open("rb") as handle:
        while True:
            try:
                ordinal, location, content = pickle.load(handle)
            except EOFError:
                break
            records.append((int(ordinal), str(location), str(content)))
    return records, len(records)


def _extract(
    protocol: _PersistentProtocol,
    request_id: int,
    source: Path,
    output: Path,
) -> tuple[int, list[tuple[int, str, str]], list[str]]:
    protocol.send(
        {
            "type": "extract",
            "request_id": request_id,
            "source": str(source),
            "output": str(output),
        }
    )
    progress_locations: list[str] = []
    while True:
        message = protocol.receive()
        message_type = message.get("type")
        if message_type == "progress":
            if message.get("request_id") != request_id:
                raise RuntimeError(f"progress request mismatch: {message!r}")
            progress_locations.append(str(message.get("location", "")))
            continue
        if message.get("request_id") != request_id:
            raise RuntimeError(f"response request mismatch: {message!r}")
        if message_type == "error":
            raise RuntimeError(
                f"frozen worker extraction failed: {message.get('error_type')}: "
                f"{message.get('error')}"
            )
        if message_type != "result":
            raise RuntimeError(f"unexpected extraction response: {message!r}")
        chunk_count = int(message.get("chunk_count", -1))
        returned_output = Path(str(message.get("output", "")))
        if os.path.normcase(os.path.abspath(returned_output)) != os.path.normcase(
            os.path.abspath(output)
        ):
            raise RuntimeError(
                f"frozen worker returned unexpected output: {returned_output!s}"
            )
        records, actual_count = _read_spool(output)
        if chunk_count != actual_count or chunk_count <= 0:
            raise RuntimeError(
                f"frozen worker spool count mismatch: result={chunk_count}, "
                f"actual={actual_count}"
            )
        return chunk_count, records, progress_locations


def _create_corpus(root: Path) -> dict[str, str]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    import pymupdf

    root.mkdir(parents=True, exist_ok=True)

    docx_path = root / "portable.docx"
    document = Document()
    document.add_paragraph("portable persistent DOCX marker")
    document.save(docx_path)

    xlsx_path = root / "portable.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "工作表 测试"
    worksheet["A1"] = "冻结 worker XLSX marker"
    worksheet["B1"] = "中文单元格"
    for row in range(2, 1_201):
        worksheet.cell(row=row, column=1, value=f"row-{row}")
    workbook.save(xlsx_path)
    workbook.close()

    pptx_path = root / "portable.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(914400, 914400, 5486400, 914400)
    textbox.text = "portable persistent PPTX marker"
    presentation.save(pptx_path)

    pdf_path = root / "portable.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "portable persistent PDF marker")
    pdf.save(pdf_path)
    pdf.close()

    return {
        "docx": "portable persistent DOCX marker",
        "xlsx": "冻结 worker XLSX marker",
        "pptx": "portable persistent PPTX marker",
        "pdf": "portable persistent PDF marker",
    }


def _matching_processes(executable: Path) -> set[int]:
    if os.name != "nt":
        return set()
    try:
        import psutil  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError("portable E2E requires psutil for orphan checking") from exc

    target = os.path.normcase(os.path.abspath(executable))
    matches: set[int] = set()
    for process in psutil.process_iter(["pid", "exe"]):
        try:
            candidate = process.info.get("exe")
            if candidate and os.path.normcase(os.path.abspath(candidate)) == target:
                matches.add(int(process.info["pid"]))
        except (psutil.NoSuchProcess, psutil.AccessDenied, OSError, TypeError, ValueError):
            continue
    return matches


def _verify_protocol(executable: Path, work_dir: Path) -> None:
    sources = _create_corpus(work_dir / "protocol-corpus")
    protocol = _PersistentProtocol(executable, work_dir)
    try:
        ready = protocol.receive()
        if ready.get("type") != "ready" or ready.get("protocol_version") != 1:
            raise RuntimeError(f"invalid frozen worker ready message: {ready!r}")
        worker_pid = int(ready.get("pid", -1))
        if worker_pid != protocol.process.pid:
            raise RuntimeError(
                f"worker PID mismatch: ready={worker_pid}, process={protocol.process.pid}"
            )

        for request_id, (name, marker) in enumerate(sources.items(), start=1):
            source = work_dir / "protocol-corpus" / f"portable.{name}"
            output = work_dir / f"protocol-chunks-{request_id}.bin"
            _count, records, progress = _extract(protocol, request_id, source, output)
            content = "\n".join(record[2] for record in records)
            if marker not in content:
                raise RuntimeError(f"{name} result does not contain its marker")
            if name == "xlsx" and not any("工作表" in location for location in progress):
                raise RuntimeError(
                    "frozen XLSX request did not forward its Unicode progress label"
                )
            output.unlink(missing_ok=True)

        if protocol.process.poll() is not None:
            raise RuntimeError("frozen worker exited before the shutdown request")
        protocol.send({"type": "shutdown"})
        shutdown = protocol.receive()
        if (
            shutdown.get("type") != "shutdown_ack"
            or shutdown.get("request_id") is not None
        ):
            raise RuntimeError(f"invalid frozen worker shutdown response: {shutdown!r}")
        code = protocol.process.wait(timeout=10)
        if code != 0:
            raise RuntimeError(f"frozen worker exited with code {code}")
        print(f"frozen persistent protocol passed: pid={worker_pid}")
    finally:
        protocol.terminate()


def _verify_production_index(executable: Path, work_dir: Path) -> None:
    root = work_dir / "production-corpus"
    _create_corpus(root)
    database = work_dir / "production.db"
    report = work_dir / "production-report.json"
    before = _matching_processes(executable)
    env = os.environ.copy()
    env.update(
        {
            "DOCSEEK_FROZEN_INDEX_SMOKE": "1",
            "DOCSEEK_FROZEN_INDEX_ROOT": str(root),
            "DOCSEEK_FROZEN_INDEX_DB": str(database),
            "DOCSEEK_FROZEN_INDEX_REPORT": str(report),
            "DOCSEEK_SCAN_BACKEND": "rust",
        }
    )
    process = subprocess.run(
        [str(executable)],
        cwd=str(work_dir),
        env=env,
        timeout=180,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    if process.returncode != 0:
        details = report.read_text(encoding="utf-8") if report.exists() else "<no report>"
        raise RuntimeError(
            f"frozen production index smoke failed with exit code "
            f"{process.returncode}: {details}"
        )
    if not report.exists():
        raise RuntimeError("frozen production index smoke did not produce a report")
    result = json.loads(report.read_text(encoding="utf-8"))
    if result.get("ok") is not True:
        raise RuntimeError(f"frozen production smoke report is not successful: {result!r}")
    after = _matching_processes(executable)
    new_processes = after - before
    if new_processes:
        raise RuntimeError(f"frozen production smoke left orphan processes: {new_processes}")
    print(
        "frozen production index passed: "
        f"backend={result.get('scan_backend')} "
        f"worker_pid={result.get('worker_pid')} "
        f"request_pids={result.get('request_pids')}"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exe", type=Path, required=True)
    args = parser.parse_args()
    executable = args.exe.resolve()
    if not executable.is_file():
        raise SystemExit(f"portable executable not found: {executable}")

    with tempfile.TemporaryDirectory(prefix="docseek-frozen-e2e-") as temp_dir:
        work_dir = Path(temp_dir)
        before = _matching_processes(executable)
        _verify_protocol(executable, work_dir)
        if _matching_processes(executable) - before:
            raise RuntimeError("frozen protocol smoke left an orphan DocSeek.exe")
        _verify_production_index(executable, work_dir)

    print("frozen persistent worker E2E passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
