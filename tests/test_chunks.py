from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from docseek.chunks import TEXT_ENCODING_SAMPLE_BYTES, _iter_text_chunks, iter_document_chunks


class _NoRewindBytesIO(io.BytesIO):
    def seek(self, *args, **kwargs):
        raise AssertionError("small text fast path must not rewind and reread the probe")


class _CountingBinaryPath:
    def __init__(self, payload: bytes, *, forbid_seek: bool = False) -> None:
        self.payload = payload
        self.forbid_seek = forbid_seek
        self.open_count = 0

    def open(self, mode: str):
        self.open_count += 1
        if mode != "rb":
            raise AssertionError(f"unexpected mode: {mode}")
        stream_type = _NoRewindBytesIO if self.forbid_seek else io.BytesIO
        return stream_type(self.payload)


class DocumentChunkExtractionTests(unittest.TestCase):
    def test_small_utf8_text_reuses_encoding_probe_without_reopen_or_rewind(self) -> None:
        path = _CountingBinaryPath(
            "第一行 客户经理\r\n第二行 信贷\n".encode("utf-8"),
            forbid_seek=True,
        )

        chunks = list(_iter_text_chunks(path, target_chars=12_000))  # type: ignore[arg-type]

        self.assertEqual(path.open_count, 1)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].location, "行 1-2")
        self.assertEqual(chunks[0].content, "第一行 客户经理\n第二行 信贷")

    def test_small_gb18030_text_reuses_probe_and_preserves_content(self) -> None:
        path = _CountingBinaryPath(
            "客户编号,姓名\r\n001,张三\r\n".encode("gb18030"),
            forbid_seek=True,
        )

        chunks = list(_iter_text_chunks(path, target_chars=12_000))  # type: ignore[arg-type]

        self.assertEqual(path.open_count, 1)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].content, "客户编号,姓名\n001,张三")

    def test_large_text_keeps_single_handle_streaming_path(self) -> None:
        payload = (
            b"A" * TEXT_ENCODING_SAMPLE_BYTES
            + b"\n"
            + "客户经理\n信贷业务\n".encode("utf-8")
        )
        path = _CountingBinaryPath(payload)

        chunks = list(_iter_text_chunks(path, target_chars=len(payload) + 100))  # type: ignore[arg-type]

        self.assertEqual(path.open_count, 1)
        self.assertEqual(len(chunks), 1)
        self.assertTrue(chunks[0].content.endswith("客户经理\n信贷业务"))

    def test_docx_tables_keep_body_order_and_heading_ownership(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "ordered.docx"
            document = Document()
            document.add_heading("第一节", level=1)
            document.add_paragraph("前文")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "中间表格"
            document.add_paragraph("后文")
            document.add_heading("第二节", level=1)
            document.add_paragraph("另一节")
            document.save(path)
            chunks = list(iter_document_chunks(path))
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0].content, "第一节\n前文\n中间表格\n后文")
        self.assertIn("第一节", chunks[0].location)
        self.assertIn("第二节", chunks[1].location)
        self.assertNotIn("中间表格", chunks[1].content)

    def test_docx_heading_is_preserved_in_chunk_location(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "制度.docx"
            document = Document()
            document.add_heading("客户经理考核办法", level=1)
            document.add_paragraph("本章节介绍客户经理考核与管理要求。")
            document.save(path)

            chunks = list(iter_document_chunks(path))

        self.assertEqual(len(chunks), 1)
        self.assertIn("· 标题 客户经理考核办法", chunks[0].location)
        self.assertIn("本章节介绍客户经理考核与管理要求", chunks[0].content)

    def test_pptx_title_is_preserved_in_slide_location(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "培训.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "风险管理"
            slide.placeholders[1].text = "客户经理风险管理培训内容"
            presentation.save(path)

            chunks = list(iter_document_chunks(path))

        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].location, "幻灯片 1 · 标题 风险管理")
        self.assertIn("风险管理", chunks[0].content)

    def test_xlsx_preserves_interior_empty_cells_and_sheet_name(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "客户清单.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "客户明细"
            worksheet.append(["客户号", "姓名", "风险等级"])
            worksheet.append(["001", None, "高风险"])
            worksheet.append([None, None, None])
            worksheet.append([None, "张三", "正常"])
            workbook.save(path)
            workbook.close()

            chunks = list(iter_document_chunks(path, xlsx_rows_per_chunk=20))

        self.assertEqual(len(chunks), 1)
        chunk = chunks[0]
        self.assertEqual(chunk.location, "工作表 客户明细 · 行 1-4")
        self.assertIn("工作表: 客户明细", chunk.content)
        self.assertIn("客户号\t姓名\t风险等级", chunk.content)
        self.assertIn("001\t\t高风险", chunk.content)
        self.assertIn("\n\t张三\t正常", chunk.content)
        self.assertNotIn("\n\n\n", chunk.content)

    def test_xlsx_repeats_sheet_context_for_each_chunk(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "分块.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "逾期清单"
            worksheet.append(["客户", "金额"])
            worksheet.append(["A", 100])
            worksheet.append(["B", 200])
            workbook.save(path)
            workbook.close()

            chunks = list(iter_document_chunks(path, xlsx_rows_per_chunk=1))

        self.assertEqual(len(chunks), 3)
        self.assertTrue(all(chunk.content.startswith("工作表: 逾期清单\n") for chunk in chunks))
        self.assertEqual(
            [chunk.location for chunk in chunks],
            [
                "工作表 逾期清单 · 行 1-1",
                "工作表 逾期清单 · 行 2-2",
                "工作表 逾期清单 · 行 3-3",
            ],
        )

    def test_xlsx_progress_is_throttled_and_reports_final_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "大表.xlsx"
            workbook = Workbook(write_only=True)
            worksheet = workbook.create_sheet("客户明细")
            for row_no in range(1, 2_506):
                worksheet.append([row_no, f"客户{row_no}"])
            workbook.save(path)
            workbook.close()

            progress: list[tuple[str, int]] = []
            list(
                iter_document_chunks(
                    path,
                    xlsx_rows_per_chunk=200,
                    on_progress=lambda location, current: progress.append((location, current)),
                )
            )

        self.assertEqual(
            progress,
            [
                ("工作表 客户明细", 1_000),
                ("工作表 客户明细", 2_000),
                ("工作表 客户明细", 2_505),
            ],
        )


if __name__ == "__main__":
    unittest.main()
