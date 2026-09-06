from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime, time
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from docseek.chunks import _iter_xlsx_chunks
from docseek.document_adapters import (
    DEFAULT_ADAPTER_REGISTRY,
    CalamineSpreadsheetAdapter,
    XlsxCalamineFastAdapter,
)


@unittest.skipUnless(
    CalamineSpreadsheetAdapter().is_available(),
    "python-calamine optional dependency is not installed",
)
class CalamineSpreadsheetAdapterTests(unittest.TestCase):
    def test_streams_rows_and_preserves_interior_empty_cells(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "客户清单.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "客户明细"
            sheet.append(["客户号", "姓名", "风险等级"])
            sheet.append(["001", None, "高风险"])
            workbook.save(path)
            workbook.close()

            progress: list[tuple[str, int]] = []
            adapter = CalamineSpreadsheetAdapter(rows_per_chunk=200)
            chunks = list(
                adapter.iter_chunks(
                    path,
                    on_progress=lambda label, row: progress.append((label, row)),
                )
            )

        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].location, "工作表 客户明细 · 行 1-2")
        self.assertIn("客户号\t姓名\t风险等级", chunks[0].content)
        self.assertIn("001\t\t高风险", chunks[0].content)
        self.assertEqual(progress[-1], ("工作表 客户明细", 2))

    def test_xlsx_fast_path_matches_openpyxl_for_common_cell_types(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "语义回归.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "类型测试"
            sheet.append(["文本", "整数", "小数", "布尔", "空值", "日期", "日期时间", "时间"])
            sheet.append(
                [
                    "客户经理",
                    100,
                    12.5,
                    True,
                    None,
                    date(2026, 9, 6),
                    datetime(2026, 9, 6, 8, 30, 15),
                    time(14, 5, 9),
                ]
            )
            sheet.append(["尾列保留", None, None, False, "末尾", None, None, None])
            second = workbook.create_sheet("第二工作表")
            second.append(["客户号", "金额"])
            second.append(["0001", 200.0])
            workbook.save(path)
            workbook.close()

            expected = list(_iter_xlsx_chunks(path, rows_per_chunk=2))
            actual = list(
                XlsxCalamineFastAdapter().iter_chunks(
                    path,
                    spreadsheet_rows_per_chunk=2,
                )
            )

        self.assertEqual(actual, expected)

    def test_xlsx_fast_path_falls_back_without_partial_chunks(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "回退.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "客户明细"
            sheet.append(["客户号", "姓名"])
            sheet.append(["001", "张三"])
            workbook.save(path)
            workbook.close()

            expected = list(_iter_xlsx_chunks(path, rows_per_chunk=200))
            with patch(
                "docseek.document_adapters._iter_calamine_spreadsheet_chunks",
                side_effect=RuntimeError("synthetic calamine failure"),
            ):
                actual = list(XlsxCalamineFastAdapter().iter_chunks(path))

        self.assertEqual(actual, expected)

    def test_default_registry_prefers_xlsx_fast_path(self) -> None:
        self.assertEqual(
            DEFAULT_ADAPTER_REGISTRY.adapter_for(Path("客户清单.xlsx")).name,
            "calamine-xlsx-fast",
        )


if __name__ == "__main__":
    unittest.main()
