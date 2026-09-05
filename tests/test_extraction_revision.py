from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from docx import Document

from docseek.chunk_store import ChunkStore
from docseek.chunks import DocumentChunk
from docseek.extraction_revision import current_extraction_revision
from docseek.indexer import DirectoryIndexer
from docseek.schema import CURRENT_SCHEMA_VERSION
from docseek.search_db import SearchDatabase


class ExtractionRevisionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name) / "root"
        self.root.mkdir()
        self.db_path = Path(self.temp_dir.name) / "docseek.db"
        self.db = SearchDatabase(self.db_path)
        self.store = ChunkStore(self.db_path)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_v7_upgrade_seeds_existing_files_at_legacy_revision_without_rebuild(self) -> None:
        path = r"C:\docs\legacy.pdf"
        self.store.replace_document(
            path=path,
            filename="legacy.pdf",
            extension=".pdf",
            modified_time=1.0,
            size=100,
            chunks=[DocumentChunk(0, "第 1 页", "信贷业务")],
        )
        with self.store.connect() as conn:
            chunk_id = int(conn.execute("SELECT id FROM chunks").fetchone()[0])
            conn.execute("DROP TABLE extraction_state")
            conn.execute("PRAGMA user_version = 7")

        migrated = ChunkStore(self.db_path)
        with migrated.connect() as conn:
            version = int(conn.execute("PRAGMA user_version").fetchone()[0])
            revision = int(
                conn.execute(
                    "SELECT revision FROM extraction_state WHERE path = ?", (path,)
                ).fetchone()[0]
            )
            migrated_chunk_id = int(conn.execute("SELECT id FROM chunks").fetchone()[0])

        self.assertEqual(version, CURRENT_SCHEMA_VERSION)
        self.assertEqual(revision, 1)
        self.assertEqual(migrated_chunk_id, chunk_id)
        self.assertEqual([row.filename for row in migrated.search("信贷")], ["legacy.pdf"])

    def test_full_scan_refreshes_only_outdated_format_once(self) -> None:
        docx_path = self.root / "制度.docx"
        document = Document()
        document.add_heading("客户经理管理", level=1)
        document.add_paragraph("信贷业务操作要求")
        document.save(docx_path)

        txt_path = self.root / "notes.txt"
        txt_path.write_text("客户经理普通笔记", encoding="utf-8")

        first = DirectoryIndexer(self.db).scan(self.root)
        self.assertEqual(first.indexed, 2)

        normalized_docx = str(docx_path.resolve())
        normalized_txt = str(txt_path.resolve())
        with self.store.connect() as conn:
            docx_revision = int(
                conn.execute(
                    "SELECT revision FROM extraction_state WHERE path = ?",
                    (normalized_docx,),
                ).fetchone()[0]
            )
            txt_revision = int(
                conn.execute(
                    "SELECT revision FROM extraction_state WHERE path = ?",
                    (normalized_txt,),
                ).fetchone()[0]
            )
            conn.execute(
                "UPDATE extraction_state SET revision = 1 WHERE path = ?",
                (normalized_docx,),
            )

        self.assertEqual(docx_revision, current_extraction_revision(".docx"))
        self.assertEqual(txt_revision, current_extraction_revision(".txt"))
        self.assertGreater(docx_revision, txt_revision)

        second = DirectoryIndexer(self.db).scan(self.root)
        self.assertEqual(second.indexed, 1)
        self.assertEqual(second.unchanged, 1)

        with self.store.connect() as conn:
            refreshed_revision = int(
                conn.execute(
                    "SELECT revision FROM extraction_state WHERE path = ?",
                    (normalized_docx,),
                ).fetchone()[0]
            )
            locations = [
                str(row[0])
                for row in conn.execute(
                    """
                    SELECT c.location
                    FROM chunks c
                    JOIN files f ON f.id = c.file_id
                    WHERE f.path = ?
                    ORDER BY c.ordinal
                    """,
                    (normalized_docx,),
                ).fetchall()
            ]

        self.assertEqual(refreshed_revision, current_extraction_revision(".docx"))
        self.assertTrue(any("标题 客户经理管理" in location for location in locations))

        third = DirectoryIndexer(self.db).scan(self.root)
        self.assertEqual(third.indexed, 0)
        self.assertEqual(third.unchanged, 2)

    def test_precise_update_refreshes_outdated_revision_even_when_file_is_unchanged(self) -> None:
        target = self.root / "guide.txt"
        target.write_text("信贷业务操作指引", encoding="utf-8")
        DirectoryIndexer(self.db).scan(self.root)

        normalized = str(target.resolve())
        with self.store.connect() as conn:
            conn.execute(
                "UPDATE extraction_state SET revision = 0 WHERE path = ?",
                (normalized,),
            )

        stats = DirectoryIndexer(self.db).update_paths([target])
        self.assertEqual(stats.indexed, 1)
        self.assertEqual(stats.unchanged, 0)

        with self.store.connect() as conn:
            revision = int(
                conn.execute(
                    "SELECT revision FROM extraction_state WHERE path = ?",
                    (normalized,),
                ).fetchone()[0]
            )
        self.assertEqual(revision, current_extraction_revision(".txt"))


if __name__ == "__main__":
    unittest.main()
